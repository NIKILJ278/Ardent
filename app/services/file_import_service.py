"""Report uploads, for platforms that publish no data API.

Myntra, Nykaa, AJIO and Meesho give sellers order reports to download from their
partner portals but no API to pull them; PhonePe and Paytm give transaction
reports the same way. This reads those files — CSV or Excel — and writes them
through the same order and payment writers as the API connectors.

Column names differ between platforms and change without notice, so columns are
matched by meaning against a list of known names. When a required column cannot
be found, the import stops and says which columns it saw, and the dashboard
lets the user point each field at the right one.
"""
import csv
import io
import re
from collections import OrderedDict
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from app.extensions import db
from app.services.base_connector import BaseConnector
from app.services.writers import catalogue_for, upsert_order, upsert_payment

IST = ZoneInfo("Asia/Kolkata")
MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS = 200_000


class ImportError_(Exception):
    """The file could not be read as a report of this kind."""

    def __init__(self, message, *, headers=None, missing=None):
        super().__init__(message)
        self.headers = headers or []
        self.missing = missing or []


# ── What each kind of report must say, and the names it might use ────────────

ORDER_FIELDS = OrderedDict([
    ("order_id", {"required": True, "label": "Order ID", "aliases": [
        "order id", "order no", "order number", "seller order id", "order release id",
        "sub order no", "suborder no", "sub order id", "store order id", "po number", "order code"]}),
    ("order_date", {"required": True, "label": "Order date", "aliases": [
        "order date", "created on", "order created date", "ordered on", "order placed date",
        "order date time", "purchase date", "date"]}),
    ("sku", {"required": True, "label": "SKU", "aliases": [
        "seller sku", "seller sku code", "sku", "sku code", "sku id", "vendor sku", "style id",
        "product id", "item sku", "article number"]}),
    ("quantity", {"required": False, "label": "Quantity", "aliases": [
        "quantity", "qty", "units", "item quantity", "quantity ordered"]}),
    ("amount", {"required": True, "label": "Item amount", "aliases": [
        "final amount", "selling price", "item price", "total price", "invoice amount",
        "order value", "amount", "seller price", "supplier discounted price",
        "supplier discounted price incl gst and commision", "listing price", "sale amount",
        "gross amount"]}),
    ("product_name", {"required": False, "label": "Product name", "aliases": [
        "product name", "product title", "item name", "product", "title", "style name", "description"]}),
    ("discount", {"required": False, "label": "Seller discount", "aliases": [
        "seller discount", "discount", "coupon discount", "discount amount"]}),
    ("status", {"required": False, "label": "Status", "aliases": [
        "order status", "status", "reason for credit entry", "item status", "shipment status"]}),
    ("returned", {"required": False, "label": "Returned amount", "aliases": [
        "return amount", "returned amount", "refund amount", "customer return amount"]}),
    ("commission", {"required": False, "label": "Marketplace fees", "aliases": [
        "commission", "total commission", "marketplace fee", "platform fee", "fees",
        "commission amount", "total fees"]}),
    ("state", {"required": False, "label": "Customer state", "aliases": [
        "customer state", "state", "shipping state", "delivery state", "ship to state"]}),
    ("category", {"required": False, "label": "Category", "aliases": [
        "category", "article type", "product category", "sub category"]}),
])

PAYMENT_FIELDS = OrderedDict([
    ("transaction_id", {"required": True, "label": "Transaction ID", "aliases": [
        "transaction id", "txn id", "merchant transaction id", "payment id", "provider reference id",
        "order id", "transaction reference"]}),
    ("date", {"required": True, "label": "Date", "aliases": [
        "transaction date", "date", "created at", "payment date", "txn date", "transaction time"]}),
    ("amount", {"required": True, "label": "Amount", "aliases": [
        "amount", "transaction amount", "txn amount", "gross amount", "total amount"]}),
    ("fee", {"required": False, "label": "Fee", "aliases": [
        "fee", "fees", "commission", "mdr", "pg charges", "total fee", "charges", "transaction fee"]}),
    ("tax", {"required": False, "label": "Tax on fee", "aliases": ["gst", "tax", "gst on fee", "igst"]}),
    ("status", {"required": False, "label": "Status", "aliases": [
        "status", "transaction status", "payment status"]}),
    ("method", {"required": False, "label": "Payment method", "aliases": [
        "payment mode", "instrument type", "payment method", "mode", "payment instrument"]}),
    ("order_ref", {"required": False, "label": "Order reference", "aliases": [
        "merchant order id", "order reference", "order id"]}),
    ("settlement_ref", {"required": False, "label": "Settlement / UTR", "aliases": [
        "utr", "settlement id", "utr no", "settlement utr", "bank reference"]}),
    ("refunded", {"required": False, "label": "Refunded amount", "aliases": [
        "refund amount", "refunded amount", "reversal amount"]}),
])

KINDS = {"orders": ORDER_FIELDS, "payments": PAYMENT_FIELDS}


def _norm(text):
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def detect(headers, kind, override=None):
    """Map each field to a column. Returns (mapping, missing_required)."""
    fields = KINDS[kind]
    by_norm = {_norm(h): h for h in headers if h}
    mapping = {}
    override = {k: v for k, v in (override or {}).items() if v}
    for key, spec in fields.items():
        if key in override:
            if override[key] not in headers:
                raise ImportError_(f"Column '{override[key]}' is not in this file", headers=headers)
            mapping[key] = override[key]
            continue
        for alias in spec["aliases"]:
            if _norm(alias) in by_norm:
                mapping[key] = by_norm[_norm(alias)]
                break
    missing = [k for k, spec in fields.items() if spec["required"] and k not in mapping]
    return mapping, missing


# ── Reading values ─────────────────────────────────────────────────────────────

def money(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text or text.lower() in {"na", "n/a", "-", "nil", "none"}:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = re.sub(r"(?i)(inr|rs\.?|₹)", "", text).replace(",", "").strip("() ")
    try:
        number = float(text)
    except ValueError:
        return None
    return -number if negative else number


_DATE_FORMATS = (
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d",
    "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d-%m-%Y",
    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
    "%d %b %Y %H:%M:%S", "%d %b %Y", "%b %d, %Y", "%d-%b-%Y", "%d-%b-%y", "%Y/%m/%d",
)


def when(value):
    """A report date, read as India time unless it says otherwise.

    Day-first formats are tried before month-first because every platform here
    is Indian; an ambiguous 03/04/2026 is 3 April.
    """
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        dt = value
    else:
        text = str(value).strip()
        dt = None
        try:
            dt = datetime.fromisoformat(text.replace("Z", "+00:00"))
        except ValueError:
            for fmt in _DATE_FORMATS:
                try:
                    dt = datetime.strptime(text, fmt)
                    break
                except ValueError:
                    continue
        if dt is None:
            return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=IST)
    return dt.astimezone(timezone.utc)


def read_table(filename, data):
    """Headers and rows from a CSV or Excel file held in memory."""
    if len(data) > MAX_BYTES:
        raise ImportError_("The file is larger than 25 MB. Split the report by date and upload each part.")
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - dependency is in requirements
            raise ImportError_("Excel support is not installed on the server; upload a CSV") from exc
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        # Reports often open with a title block; the header is the first row
        # with at least three filled cells.
        headers = None
        for row in rows:
            filled = [c for c in row if c not in (None, "")]
            if len(filled) >= 3:
                headers = [str(c).strip() if c is not None else "" for c in row]
                break
        if headers is None:
            raise ImportError_("No header row was found in the first sheet")
        body = [dict(zip(headers, r)) for r in rows if any(c not in (None, "") for c in r)]
    elif name.endswith((".csv", ".txt")):
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        headers = [h.strip() for h in (reader.fieldnames or [])]
        reader.fieldnames = headers
        body = [r for r in reader if any((v or "").strip() for v in r.values() if isinstance(v, str))]
    elif name.endswith(".xls"):
        raise ImportError_("Old .xls files are not supported. Open it and save as .xlsx or .csv")
    else:
        raise ImportError_("Upload a .csv or .xlsx report")
    if len(body) > MAX_ROWS:
        raise ImportError_(f"The report has more than {MAX_ROWS:,} rows. Split it by date.")
    return headers, body


# ── Connector ────────────────────────────────────────────────────────────────

class FileImportConnector(BaseConnector):
    """A source whose data arrives as uploaded reports.

    `kind` is "orders" for marketplaces and "payments" for gateways.
    """

    platform_name = "file"
    kind = "orders"

    def test_connection(self):
        return {"account_id": self.connection.platform, "display_name": self.connection.display_name}

    def sync(self, since=None):
        return {"platform": self.connection.platform, "records_synced": 0, "message":
                "This source has no API. Upload a new report to update it."}

    def import_file(self, filename, data, override=None):
        headers, rows = read_table(filename, data)
        mapping, missing = detect(headers, self.kind, override)
        if missing:
            labels = [KINDS[self.kind][k]["label"] for k in missing]
            raise ImportError_(
                f"Could not find: {', '.join(labels)}. Choose which column holds each.",
                headers=headers, missing=missing,
            )
        result = self._orders(rows, mapping) if self.kind == "orders" else self._payments(rows, mapping)
        result.update({"file": filename, "rows": len(rows), "mapping": mapping})
        self.mark_synced()
        db.session.commit()
        return result

    def _orders(self, rows, mapping):
        get = lambda row, key: row.get(mapping[key]) if key in mapping else None  # noqa: E731
        brand_id = self.connection.brand_id
        channel = self.connection.platform
        grouped = OrderedDict()
        skipped = 0

        for row in rows:
            order_id = str(get(row, "order_id") or "").strip()
            sku = str(get(row, "sku") or "").strip()
            amount = money(get(row, "amount"))
            if not order_id or not sku or amount is None:
                skipped += 1
                continue
            qty = int(money(get(row, "quantity")) or 1)
            status = str(get(row, "status") or "").strip().lower()
            entry = grouped.setdefault(order_id, {
                "date": when(get(row, "order_date")), "lines": [], "statuses": set(),
                "fee": None, "state": get(row, "state"), "returns": 0.0,
            })
            entry["statuses"].add(status)
            fee = money(get(row, "commission"))
            if fee is not None:
                entry["fee"] = (entry["fee"] or 0.0) + abs(fee)
            returned = abs(money(get(row, "returned")) or 0.0)
            if not returned and re.search(r"return|rto|refund", status):
                # A row marked returned with no amount returns the whole line.
                returned = amount
            entry["returns"] += returned

            known = catalogue_for(brand_id, sku)
            entry["lines"].append({
                "external_line_id": f"{order_id}:{sku}:{len(entry['lines'])}",
                "sku": sku,
                "product_name": known.get("product_name") or get(row, "product_name") or sku,
                "product_external_id": known.get("product_external_id") or f"{channel}:{sku}",
                "variant_external_id": known.get("variant_external_id") or f"sku:{sku}",
                "variant_title": known.get("variant_title"),
                "category": known.get("category") or get(row, "category"),
                "subcategory": known.get("subcategory"),
                "quantity": qty,
                # Reports state the line value; the per-unit price is derived.
                "unit_price": amount / qty if qty else amount,
                "unit_cost": known.get("unit_cost"),
                "discount": abs(money(get(row, "discount")) or 0.0),
                "returned_amount": returned,
                "returned_quantity": qty if returned else 0,
            })

        written = 0
        for order_id, entry in grouped.items():
            statuses = entry["statuses"]
            cancelled = bool(statuses) and all("cancel" in s for s in statuses)
            rto = any("rto" in s for s in statuses)
            upsert_order(
                brand_id, channel, order_id,
                order_name=order_id,
                order_date=entry["date"] or datetime.now(timezone.utc),
                lines=entry["lines"],
                currency="INR",
                status="delivered" if any("deliver" in s for s in statuses) else "placed",
                is_cancelled=cancelled,
                is_rto=rto,
                shipping_state=entry["state"],
                marketplace_fee=entry["fee"],
                returns=[(entry["returns"], None, "rto" if rto else "return", None)] if entry["returns"] else (),
            )
            written += 1
        db.session.commit()
        return {"platform": channel, "orders": written, "skipped_rows": skipped, "records_synced": written}

    def _payments(self, rows, mapping):
        get = lambda row, key: row.get(mapping[key]) if key in mapping else None  # noqa: E731
        brand_id = self.connection.brand_id
        gateway = self.connection.platform
        written = skipped = 0
        for row in rows:
            txn = str(get(row, "transaction_id") or "").strip()
            amount = money(get(row, "amount"))
            if not txn or amount is None:
                skipped += 1
                continue
            fee = money(get(row, "fee"))
            upsert_payment(
                brand_id, gateway, txn,
                order_ref=(str(get(row, "order_ref")).strip() or None) if get(row, "order_ref") else None,
                status=(str(get(row, "status") or "").strip().lower() or None),
                method=(str(get(row, "method") or "").strip().lower() or None),
                currency="INR",
                amount=abs(amount),
                fee=abs(fee) if fee is not None else None,
                tax=abs(money(get(row, "tax"))) if money(get(row, "tax")) is not None else None,
                refunded=abs(money(get(row, "refunded")) or 0.0),
                settlement_ref=(str(get(row, "settlement_ref")).strip() or None) if get(row, "settlement_ref") else None,
                occurred_at=when(get(row, "date")) or datetime.now(timezone.utc),
            )
            written += 1
        db.session.commit()
        return {"platform": gateway, "payments": written, "skipped_rows": skipped, "records_synced": written}


class OrderReportConnector(FileImportConnector):
    kind = "orders"


class PaymentReportConnector(FileImportConnector):
    kind = "payments"
